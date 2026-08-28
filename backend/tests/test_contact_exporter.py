from datetime import UTC, date, datetime
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app.modules.contacts.exporter import export_new_contacts_xlsx
from app.modules.contacts.schemas import NewContactItem


def test_new_contact_export_contains_typed_rows_and_filters() -> None:
    content = export_new_contacts_xlsx(
        [
            NewContactItem(
                id=uuid4(),
                phone_number="+77012345678",
                first_contact_at=datetime(2026, 8, 28, 9, 30, tzinfo=UTC),
                source="whatsapp",
                last_contact_at=datetime(2026, 8, 28, 10, 15, tzinfo=UTC),
                inbound_count=3,
                call_count=1,
                message_count=2,
            )
        ],
        date(2026, 8, 1),
        date(2026, 8, 31),
    )

    workbook = load_workbook(BytesIO(content))
    sheet = workbook["Новые обращения"]

    assert sheet["A1"].value == "Новые обращения"
    assert sheet["A6"].value == "+77012345678"
    assert isinstance(sheet["B6"].value, datetime)
    assert sheet["C6"].value == "WhatsApp"
    assert sheet["E6"].value == 3
    assert sheet["H6"].value == "Не найден в 1С"
    assert sheet.freeze_panes == "A6"
    assert sheet.tables["NewContacts"].ref == "A5:H6"
