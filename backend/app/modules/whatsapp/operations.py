"""Operational monitoring, reporting and exports for the WhatsApp channel."""

import base64
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import Settings
from app.core.errors import AppError
from app.modules.auth.models import User, UserRole
from app.modules.contacts.models import ContactIdentity
from app.modules.whatsapp.models import (
    WhatsAppConversation,
    WhatsAppGatewayHealth,
    WhatsAppGatewayLog,
    WhatsAppMessage,
)
from app.modules.whatsapp.schemas import (
    WhatsAppAnalyticsResponse,
    WhatsAppGatewayHeartbeat,
    WhatsAppGatewayLogItem,
    WhatsAppGatewayLogResponse,
    WhatsAppQrStatusResponse,
)
from app.modules.whatsapp.security import decrypt_contact, mask_contact
from app.shared.timezone import clinic_day_end_exclusive, clinic_day_start


def _moment(value: int | None) -> datetime | None:
    try:
        return datetime.fromtimestamp(int(value), tz=UTC) if value is not None else None
    except (TypeError, ValueError, OSError):
        return None


class WhatsAppOperationsService:
    def __init__(self, session: AsyncSession, settings: Settings) -> None:
        self.session = session
        self.settings = settings

    async def record_heartbeat(
        self, tenant_id: UUID, payload: WhatsAppGatewayHeartbeat
    ) -> None:
        now = datetime.now(UTC)
        health = await self.session.scalar(
            select(WhatsAppGatewayHealth).where(
                WhatsAppGatewayHealth.tenant_id == tenant_id
            )
        )
        values = {
            "state": payload.state,
            "connected": payload.connected,
            "phone_masked": mask_contact(payload.phone) if payload.phone else None,
            "gateway_version": payload.gateway_version,
            "instance_id": payload.instance_id,
            "started_at": _moment(payload.started_at),
            "last_heartbeat_at": now,
            "last_message_at": _moment(payload.last_message_at),
            "last_history_sync_at": _moment(payload.last_history_sync_at),
            "messages_forwarded": payload.messages_forwarded,
            "history_messages_forwarded": payload.history_messages_forwarded,
            "reconnect_count": payload.reconnect_count,
            "last_error": payload.last_error,
        }
        if health is None:
            health = WhatsAppGatewayHealth(tenant_id=tenant_id, **values)
            self.session.add(health)
        else:
            for key, value in values.items():
                setattr(health, key, value)
        for item in payload.logs:
            self.session.add(
                WhatsAppGatewayLog(
                    tenant_id=tenant_id,
                    level=item.level,
                    event=item.event,
                    message=item.message,
                    occurred_at=_moment(item.timestamp) or now,
                )
            )
        # Operational logs are intentionally short lived and contain no message bodies.
        await self.session.execute(
            delete(WhatsAppGatewayLog).where(
                WhatsAppGatewayLog.tenant_id == tenant_id,
                WhatsAppGatewayLog.occurred_at < now - timedelta(days=30),
            )
        )

    async def cached_status(self, user: User) -> WhatsAppQrStatusResponse:
        self._allowed(user)
        health = await self.session.scalar(
            select(WhatsAppGatewayHealth).where(
                WhatsAppGatewayHealth.tenant_id == user.tenant_id
            )
        )
        configured = bool(
            self.settings.whatsapp_qr_gateway_url
            and self.settings.whatsapp_qr_gateway_secret.get_secret_value()
        )
        if health is None:
            return WhatsAppQrStatusResponse(
                configured=configured,
                state="unknown" if configured else "not_configured",
                connected=False,
                message=(
                    "Шлюз ещё не присылал данные о состоянии"
                    if configured
                    else "QR-шлюз ещё не развёрнут"
                ),
                stale=True,
            )
        stale = datetime.now(UTC) - health.last_heartbeat_at > timedelta(minutes=2)
        return WhatsAppQrStatusResponse(
            configured=configured,
            state="offline" if stale else health.state,
            connected=health.connected and not stale,
            phone=health.phone_masked,
            message=("Нет heartbeat больше двух минут" if stale else "Состояние получено от Baileys"),
            last_heartbeat_at=health.last_heartbeat_at,
            last_message_at=health.last_message_at,
            last_history_sync_at=health.last_history_sync_at,
            messages_forwarded=health.messages_forwarded,
            history_messages_forwarded=health.history_messages_forwarded,
            reconnect_count=health.reconnect_count,
            last_error=health.last_error,
            stale=stale,
        )

    async def logs(self, user: User, limit: int = 100) -> WhatsAppGatewayLogResponse:
        self._allowed(user)
        rows = list(
            (
                await self.session.scalars(
                    select(WhatsAppGatewayLog)
                    .where(WhatsAppGatewayLog.tenant_id == user.tenant_id)
                    .order_by(WhatsAppGatewayLog.occurred_at.desc())
                    .limit(limit)
                )
            ).all()
        )
        return WhatsAppGatewayLogResponse(
            items=[
                WhatsAppGatewayLogItem(
                    id=row.id,
                    level=row.level,
                    event=row.event,
                    message=row.message,
                    occurred_at=row.occurred_at,
                )
                for row in rows
            ]
        )

    async def analytics(
        self, user: User, date_from: date, date_to: date
    ) -> WhatsAppAnalyticsResponse:
        self._allowed(user)
        if date_from > date_to or (date_to - date_from).days > 366:
            raise AppError("INVALID_DATE_RANGE", "Период должен быть от 1 до 367 дней", 422)
        start = clinic_day_start(date_from)
        end = clinic_day_end_exclusive(date_to)
        event_at = func.coalesce(WhatsAppMessage.provider_timestamp, WhatsAppMessage.created_at)
        rows = list(
            (
                await self.session.scalars(
                    select(WhatsAppMessage)
                    .where(
                        WhatsAppMessage.tenant_id == user.tenant_id,
                        event_at >= start,
                        event_at < end,
                    )
                    .order_by(WhatsAppMessage.conversation_id, event_at)
                    .limit(100_000)
                )
            ).all()
        )
        by_conversation: dict[UUID, list[WhatsAppMessage]] = {}
        for row in rows:
            by_conversation.setdefault(row.conversation_id, []).append(row)
        response_times: list[float] = []
        unanswered = 0
        for messages in by_conversation.values():
            if messages and messages[-1].direction == "in":
                unanswered += 1
            first_in = next((item for item in messages if item.direction == "in"), None)
            if first_in is None:
                continue
            first_out = next(
                (
                    item
                    for item in messages
                    if item.direction == "out" and self._event_at(item) >= self._event_at(first_in)
                ),
                None,
            )
            if first_out is not None:
                response_times.append(
                    max(0.0, (self._event_at(first_out) - self._event_at(first_in)).total_seconds())
                )
        conversation_ids = list(by_conversation)
        waiting = 0
        if conversation_ids:
            waiting = int(
                await self.session.scalar(
                    select(func.count())
                    .select_from(WhatsAppConversation)
                    .where(
                        WhatsAppConversation.tenant_id == user.tenant_id,
                        WhatsAppConversation.id.in_(conversation_ids),
                        WhatsAppConversation.state == "human_requested",
                    )
                )
                or 0
            )
        new_contacts = int(
            await self.session.scalar(
                select(func.count())
                .select_from(ContactIdentity)
                .where(
                    ContactIdentity.tenant_id == user.tenant_id,
                    ContactIdentity.first_inbound_source == "whatsapp",
                    ContactIdentity.first_inbound_at >= start,
                    ContactIdentity.first_inbound_at < end,
                    ContactIdentity.was_known_patient.is_(False),
                )
            )
            or 0
        )
        existing_patient_contacts = int(
            await self.session.scalar(
                select(func.count())
                .select_from(ContactIdentity)
                .where(
                    ContactIdentity.tenant_id == user.tenant_id,
                    ContactIdentity.first_inbound_source == "whatsapp",
                    ContactIdentity.first_inbound_at >= start,
                    ContactIdentity.first_inbound_at < end,
                    ContactIdentity.was_known_patient.is_(True),
                )
            )
            or 0
        )
        return WhatsAppAnalyticsResponse(
            date_from=date_from,
            date_to=date_to,
            conversations=len(by_conversation),
            messages_total=len(rows),
            incoming_messages=sum(row.direction == "in" for row in rows),
            outgoing_messages=sum(row.direction == "out" for row in rows),
            bot_messages=sum(row.sender_kind == "bot" for row in rows),
            human_messages=sum(row.sender_kind in {"human", "business_app"} for row in rows),
            history_messages=sum(row.status == "history" for row in rows),
            media_messages=sum(row.message_type != "text" for row in rows),
            voice_messages=sum(row.message_type == "audio" for row in rows),
            transcribed_voice_messages=sum(bool(row.transcript_ciphertext) for row in rows),
            waiting_for_human=waiting,
            unanswered_conversations=unanswered,
            average_first_response_seconds=(
                round(sum(response_times) / len(response_times), 1) if response_times else None
            ),
            new_contacts=new_contacts,
            existing_patient_contacts=existing_patient_contacts,
        )

    async def export_xlsx(self, user: User, date_from: date, date_to: date) -> bytes:
        self._allowed(user)
        if date_from > date_to or (date_to - date_from).days > 366:
            raise AppError("INVALID_DATE_RANGE", "Период должен быть от 1 до 367 дней", 422)
        start = clinic_day_start(date_from)
        end = clinic_day_end_exclusive(date_to)
        event_at = func.coalesce(WhatsAppMessage.provider_timestamp, WhatsAppMessage.created_at)
        rows = (
            await self.session.execute(
                select(WhatsAppMessage, WhatsAppConversation)
                .join(WhatsAppConversation, WhatsAppConversation.id == WhatsAppMessage.conversation_id)
                .where(
                    WhatsAppMessage.tenant_id == user.tenant_id,
                    event_at >= start,
                    event_at < end,
                )
                .order_by(event_at)
                .limit(100_000)
            )
        ).all()
        secret = self.settings.whatsapp_data_key.get_secret_value()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "WhatsApp"
        headers = [
            "Дата и время", "Номер", "Направление", "Отправитель", "Тип",
            "Сообщение", "Расшифровка", "Статус", "Диалог ID",
        ]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2F765F")
        for message, conversation in rows:
            try:
                phone = decrypt_contact(conversation.contact_ciphertext, secret)
                body = decrypt_contact(message.body_ciphertext, secret) if message.body_ciphertext else ""
                transcript = (
                    decrypt_contact(message.transcript_ciphertext, secret)
                    if message.transcript_ciphertext
                    else ""
                )
            except Exception as exc:
                raise AppError("WHATSAPP_DATA_KEY_INVALID", "Не удалось расшифровать сообщения", 503) from exc
            sheet.append(
                [
                    message.provider_timestamp or message.created_at,
                    self._excel_text(phone),
                    "Входящее" if message.direction == "in" else "Исходящее",
                    message.sender_kind,
                    message.message_type,
                    self._excel_text(body),
                    self._excel_text(transcript),
                    self._excel_text(message.status),
                    self._excel_text(str(message.conversation_id)),
                ]
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        widths = {"A": 21, "B": 18, "C": 14, "D": 18, "E": 15, "F": 60, "G": 60, "H": 16, "I": 38}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    async def media(self, user: User, message_id: UUID) -> tuple[bytes, str, str]:
        self._allowed(user)
        row = await self.session.scalar(
            select(WhatsAppMessage).where(
                WhatsAppMessage.tenant_id == user.tenant_id,
                WhatsAppMessage.id == message_id,
            )
        )
        if row is None or not row.media_ciphertext:
            raise AppError("WHATSAPP_MEDIA_NOT_FOUND", "Вложение не найдено", 404)
        try:
            encoded = decrypt_contact(
                row.media_ciphertext, self.settings.whatsapp_data_key.get_secret_value()
            )
            data = base64.b64decode(encoded, validate=True)
        except Exception as exc:
            raise AppError("WHATSAPP_MEDIA_INVALID", "Вложение повреждено", 503) from exc
        return data, row.media_mime_type or "application/octet-stream", row.media_filename or "attachment"

    async def retry_message(self, user: User, message_id: UUID) -> None:
        self._allowed(user)
        row = await self.session.scalar(
            select(WhatsAppMessage).where(
                WhatsAppMessage.tenant_id == user.tenant_id,
                WhatsAppMessage.id == message_id,
                WhatsAppMessage.direction == "out",
            )
        )
        if row is None:
            raise AppError("WHATSAPP_MESSAGE_NOT_FOUND", "Сообщение не найдено", 404)
        if row.status not in {"failed", "draft"}:
            raise AppError("WHATSAPP_MESSAGE_NOT_RETRYABLE", "Сообщение уже отправляется или отправлено", 409)
        row.status = "queued"
        row.is_draft = False
        row.delivery_attempts = 0
        row.last_delivery_error = None
        row.next_delivery_at = datetime.now(UTC)

    @staticmethod
    def _allowed(user: User) -> None:
        if user.role not in {
            UserRole.OWNER,
            UserRole.MANAGER,
            UserRole.ADMINISTRATOR,
            UserRole.SALES_MANAGER,
        }:
            raise AppError("FORBIDDEN", "WhatsApp is not available for this role", 403)

    @staticmethod
    def _event_at(message: WhatsAppMessage) -> datetime:
        return message.provider_timestamp or message.created_at

    @staticmethod
    def _excel_text(value: str) -> str:
        # Patient-controlled text must never become an executable Excel formula.
        return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value
